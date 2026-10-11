"""Backend del chat IA (sv-excel-agent + DeepSeek).

Funciona en dos modos:
- Local:  uvicorn api.index:app --port 8001   (desde la raiz del repo)
- Vercel: funcion serverless Python (ver vercel.json)

Diseno serverless: cada peticion es autocontenida (las funciones no guardan
estado entre invocaciones). El Excel modificado se devuelve en base64 dentro
de la misma respuesta del chat; el frontend lo convierte en descarga.
"""

import asyncio
import base64
import os
import shutil
import sys
import tempfile
import uuid
from pathlib import Path
from typing import Optional

# Los paquetes excel_agent / excel_mcp viven en la raiz del repo.
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi import FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
import dotenv

dotenv.load_dotenv()

from excel_mcp.excel_server import mcp
from excel_agent.agent_runner import ExcelAgentRunner, TaskInput
from excel_agent.config import ExperimentConfig

app = FastAPI(title="Excel Chat IA API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MCP_SERVER_HOST = "127.0.0.1"
MCP_SERVER_PORT = 8765
MCP_SERVER_URL = f"http://{MCP_SERVER_HOST}:{MCP_SERVER_PORT}/sse"

SYSTEM_ES = (
    ExperimentConfig.model_fields["system_prompt"].default
    + "\n\nIMPORTANTE: Responde SIEMPRE en espanol."
)

mcp_server_task: Optional[asyncio.Task] = None

# Celdas clave del modelo (una hoja de simulacion por obligacion).
KEY_CELLS = {
    "A8": "Obligacion", "B8": "Capital", "B4": "FDesembolso", "C8": "FTransferencia",
    "D8": "FVctoNeto", "E8": "FVctoTotal", "F8": "TasaNegocio",
    "L8": "TasaRemuneratoria", "AB8": "TasaDPP",
}


def build_overview(path: Path, max_sheets: int = 15) -> str:
    """Resumen del workbook para que el agente no pierda pasos explorando.

    Equivalente rapido (sin LLM) del `get_overview` del Apps Script original:
    por cada hoja de simulacion extrae obligacion, capital, fechas y tasas.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    lines = [f"Hojas ({len(wb.sheetnames)}): " + ", ".join(wb.sheetnames[:max_sheets])]
    protected = {"tasa usura", "obligaciones"}
    for ws in wb.worksheets[:max_sheets]:
        if ws.title.strip().lower() in protected:
            lines.append("- " + ws.title + ": hoja de referencia")
            continue
        try:
            a8, b8 = ws["A8"].value, ws["B8"].value
        except Exception:
            lines.append(f"- {ws.title}: (no legible)")
            continue
        if isinstance(a8, str) and a8.strip() and isinstance(b8, (int, float)):
            parts = []
            for cell, label in KEY_CELLS.items():
                try:
                    v = ws[cell].value
                except Exception:
                    v = None
                parts.append(f"{label}={v}")
            lines.append(f"- {ws.title}: " + ", ".join(parts))
        else:
            try:
                dim = ws.dimensions
            except Exception:
                dim = "?"
            lines.append(f"- {ws.title}: hoja base ({dim})")
    wb.close()
    return "\n".join(lines)


async def ensure_mcp_server_running():
    global mcp_server_task
    if mcp_server_task is None or mcp_server_task.done():
        mcp_server_task = asyncio.create_task(
            mcp.run_sse_async(host=MCP_SERVER_HOST, port=MCP_SERVER_PORT, log_level="warning")
        )
        await asyncio.sleep(1)


@app.get("/api/health")
async def health():
    return {"status": "ok"}


@app.post("/api/chat")
async def chat(
    message: str = Form(...),
    file: UploadFile = File(...),
    history: str = Form(default="[]"),
):
    """Procesa un mensaje del chat con el agente Excel.

    Recibe el Excel actual + la pregunta, ejecuta el agente (DeepSeek con
    herramientas MCP de Excel) y devuelve la respuesta junto con el Excel
    modificado en base64.
    """
    await ensure_mcp_server_running()

    tmpdir = Path(tempfile.mkdtemp(prefix="excelchat_"))
    session_id = f"chat_{uuid.uuid4().hex[:8]}"
    input_path = tmpdir / "input.xlsx"
    output_path = tmpdir / "output.xlsx"

    try:
        content = await file.read()
        input_path.write_bytes(content)
        shutil.copy(input_path, output_path)

        model_string = os.getenv("EXCEL_ASSISTANT_MODEL", "deepseek:deepseek-chat")
        config = ExperimentConfig(model=model_string, system_prompt=SYSTEM_ES, timeout=55.0)
        runner = ExcelAgentRunner(config=config, mcp_server_url=MCP_SERVER_URL)

        try:
            overview = build_overview(input_path)
        except Exception:
            overview = "(no se pudo pre-leer el archivo)"
        instruction = (
            "[Contexto del archivo]\n" + overview +
            "\n\nUsa este contexto para responder sin explorar de mas; "
            "solo usa herramientas para lo que el contexto no cubra.\n\n"
            "[Pregunta del usuario]\n" + message
        )
        agent_response = await runner.run_excel_agent(
            TaskInput(instruction=instruction, input_file=str(input_path), output_file=str(output_path))
        )

        if "timed out" in agent_response:
            return JSONResponse({"success": False, "error": "La consulta tardó demasiado. Prueba con una pregunta más específica o sobre una sola obligación."}, status_code=504)
        if agent_response.startswith("error:") or agent_response.startswith("Session init failed"):
            return JSONResponse({"success": False, "error": agent_response}, status_code=500)

        file_b64 = base64.b64encode(output_path.read_bytes()).decode("ascii")
        return JSONResponse(
            {
                "success": True,
                "response": agent_response,
                "file_base64": file_b64,
                "file_name": "modificado.xlsx",
                "session_id": session_id,
            }
        )
    except asyncio.TimeoutError:
        return JSONResponse({"success": False, "error": "El agente tardó demasiado (timeout)."}, status_code=504)
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
